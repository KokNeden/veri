"""Uzman formu şablonlarını üretir: etki_matrisi_sablon.csv / .xlsx ve yardımcı CSV'ler.

    python sablon_uret.py ../    # 04-siralama-yontemi/ klasörüne yazar
"""
from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

IDS = [f"S{i:02d}" for i in range(1, 24)]
ADLAR = ["Eğitimin niteliği", "Beceri uyumsuzluğu ve mesleki eğitim", "Beyin göçü",
         "Kadınların iş gücüne düşük katılımı", "Demografik dönüşüm",
         "Hukukun üstünlüğü ve yargının güvenilirliği ile hızı", "Plan-uygulama kopukluğu ve politika sürekliliği",
         "Aşırı merkeziyetçilik ve yerel kapasite", "Kamuda liyakat ve kurumsal kapasite", "Kayıt dışılık",
         "Toplumsal güven", "Medya ve bilgi ekosistemi (aday)", "Tarımda verimlilik ve parçalı arazi",
         "Sanayinin teknoloji düzeyi ve ithal ara malı bağımlılığı", "Ar-Ge'yi ürüne dönüştürme kapasitesi",
         "Plansız kentleşme", "Deprem ve afet riski yönetimi", "Enerjide dışa bağımlılık",
         "Su yönetimi ve kuraklık", "Bölgesel eşitsizlik", "Kronik enflasyon ve makroekonomik istikrarsızlık",
         "Düşük tasarruf ve dış finansman bağımlılığı", "Verimlilik artışının zayıflığı (izleniyor)"]
OKU = [
    "Kök Neden · Etki matrisi formu v0.1",
    "",
    "Her hücrede soru aynı: SATIRDAKİ sorun kendiliğinden düzelseydi, SÜTUNDAKİ sorun ne kadar düzelirdi?",
    "Yalnız DOĞRUDAN etkiyi puanlayın. Dolaylı etkileri yöntem kendisi hesaplar.",
    "0 = etki yok · 1 = zayıf · 2 = orta · 3 = güçlü · 4 = çok güçlü (ana belirleyici)",
    "3 ve 4 puan için 'Gerekçeler' sayfasına en az bir akademik veya kurumsal kaynak yazmak zorunludur.",
    "Emin olmadığınız hücreyi BOŞ bırakın. Boş hücre 'emin değilim' sayılır, 0 sayılmaz.",
    "Köşegen (bir sorunun kendine etkisi) 0'dır.",
    "Tahmini süre: 60-90 dakika. Satır satır ilerlemenizi öneririz.",
    "Doldurma kılavuzu: kokneden.org/uzman-paneli",
]


def main(hedef: Path) -> None:
    m = pd.DataFrame("", index=pd.Index(IDS, name="etkileyen"), columns=IDS, dtype=object)
    for i in IDS:
        m.loc[i, i] = "0"
    m.to_csv(hedef / "etki_matrisi_sablon.csv")
    pd.DataFrame({"sorun_id": IDS, "ad": ADLAR}).to_csv(hedef / "etki_matrisi_sablon_sorunlar.csv", index=False)
    pd.DataFrame(columns=["kaynak", "hedef", "puan", "durum", "gerekce", "kaynakca", "doi_url"]).to_csv(
        hedef / "etki_matrisi_sablon_gerekceler.csv", index=False)

    wb = Workbook()
    ws = wb.active
    ws.title = "Oku beni"
    for r in OKU:
        ws.append([r])
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 120

    ws2 = wb.create_sheet("Matris")
    ws2.append(["etkileyen ↓ / etkilenen →", ""] + IDS)
    ws2.append(["", ""] + [a[:30] for a in ADLAR])
    dv = DataValidation(type="whole", operator="between", formula1="0", formula2="4", allow_blank=True,
                        showErrorMessage=True, error="0-4 arası tamsayı girin")
    ws2.add_data_validation(dv)
    gri = PatternFill("solid", fgColor="DDDDDD")
    for k, (i, a) in enumerate(zip(IDS, ADLAR)):
        ws2.append([i, a] + [0 if j == i else None for j in IDS])
        for c in range(3, 3 + len(IDS)):
            cell = ws2.cell(row=k + 3, column=c)
            cell.alignment = Alignment(horizontal="center")
            if c - 3 == k:
                cell.fill = gri
            else:
                dv.add(cell)
    ws2.column_dimensions["B"].width = 44
    ws2.freeze_panes = "C3"
    for c in range(3, 3 + len(IDS)):
        ws2.column_dimensions[ws2.cell(row=1, column=c).column_letter].width = 5.5
    for c in ws2[2]:
        c.alignment = Alignment(text_rotation=90)
    ws2.row_dimensions[2].height = 160

    ws3 = wb.create_sheet("Gerekçeler")
    ws3.append(["kaynak (S..)", "hedef (S..)", "puan", "gerekçe (1-2 cümle)", "kaynakça (tam künye)", "DOI / URL"])
    for c in ws3[1]:
        c.font = Font(bold=True)
    for col, w in zip("ABCDEF", [12, 12, 6, 60, 70, 35]):
        ws3.column_dimensions[col].width = w

    ws4 = wb.create_sheet("Beyan")
    for r in ["Ad soyad", "Kurum / unvan", "Uzmanlık alanı (sütun)", "Çıkar çatışması beyanı (yoksa 'yok')",
              "Katkınız nasıl anılsın? (tam ad / baş harfler / anonim)", "Tarih"]:
        ws4.append([r, ""])
    ws4.column_dimensions["A"].width = 60
    ws4.column_dimensions["B"].width = 50
    wb.save(hedef / "etki_matrisi_sablon.xlsx")


if __name__ == "__main__":
    main(Path(sys.argv[1] if len(sys.argv) > 1 else "."))
